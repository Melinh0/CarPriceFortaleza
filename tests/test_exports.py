import io
import os
import time
import zipfile

from openpyxl import load_workbook

from app import create_app
from app.services import comparativo, exports, report, search

DATA_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app", "data"
)
CONFIG = {"DATA_DIR": DATA_DIR, "CACHE_DIR": os.path.join(DATA_DIR, "cache")}


def _resultado_base(**extras) -> dict:
    filtros = {"marca": "Peugeot", "modelo": "208"}
    ofertas = search.ofertas_estimadas(filtros, search.carregar_catalogo(DATA_DIR))
    resultado = {
        "id": "export00000001",
        "ts": time.time(),
        "filtros": filtros,
        "filtros_relaxados": [],
        "ofertas": ofertas,
        "estatisticas": search._estatisticas(ofertas),
        "mencoes_web": [
            {
                "titulo": "Peugeot 208 0 km à vista",
                "url": "https://exemplo.com/208",
                "trecho": "preço na concessionária",
                "preco": 89900.0,
            }
        ],
        "comentarios": {
            "resumo_base": {
                "pontos_positivos": ["econômico"],
                "pontos_negativos": ["porta-malas pequeno"],
                "sobre_preco": ["bom custo"],
            },
            "comentarios_web": [
                {
                    "titulo": "Donos falam do 208",
                    "url": "https://exemplo.com/forum",
                    "trecho": "confortável no dia a dia",
                }
            ],
            "origem": {"base_local": "base local"},
        },
        "fontes_ao_vivo": 0,
        "consulta_portais": False,
        "fontes_status": [
            {"nome": "Portais de anuncios", "ok": False, "detalhe": "nenhum"},
        ],
        "erros": [{"fonte": "portais", "motivo": "bloqueio"}],
    }
    resultado.update(extras)
    return resultado


def _relatorio() -> dict:
    return report.gerar_relatorio(_resultado_base(), CONFIG)


def _comparativo(monkeypatch) -> tuple[str, dict]:
    monkeypatch.setattr(
        comparativo,
        "pesquisar_detalhes",
        lambda marca, modelo: {"consultas": [], "mencoes": []},
    )
    return comparativo.executar_comparativo(
        [{"marca": "BYD", "modelo": "Dolphin Mini"}], {}, CONFIG
    )


def test_relatorio_pdf_gera_arquivo_valido():
    dados = exports.relatorio_pdf(_relatorio())
    assert dados.startswith(b"%PDF")
    assert len(dados) > 1000


def test_relatorio_planilha_abre_com_abas_organizadas():
    dados = exports.relatorio_xlsx(_relatorio())
    assert zipfile.is_zipfile(io.BytesIO(dados))
    pasta = load_workbook(io.BytesIO(dados))
    assert pasta.sheetnames == [
        "Resumo",
        "Ofertas",
        "Concessionárias",
        "Comparativo da marca",
        "Comentários",
        "Menções web",
    ]
    ofertas = pasta["Ofertas"]
    assert ofertas.cell(row=1, column=1).value == "Veículo"
    assert ofertas.max_row > 1
    assert pasta["Resumo"].cell(row=1, column=1).value == "Relatório"
    assert pasta["Concessionárias"].max_row > 1


def test_comparativo_gera_pdf_e_planilha(monkeypatch):
    _, rel = _comparativo(monkeypatch)
    dados = exports.comparativo_pdf(rel)
    assert dados.startswith(b"%PDF")

    planilha = exports.comparativo_xlsx(rel)
    assert zipfile.is_zipfile(io.BytesIO(planilha))
    pasta = load_workbook(io.BytesIO(planilha))
    assert pasta.sheetnames == [
        "Resumo",
        "Preços",
        "Parcelas",
        "Detalhes",
        "Concessionárias",
    ]
    assert pasta["Preços"].max_row > 1
    assert pasta["Parcelas"].cell(row=1, column=1).value == "Veículo"


def test_rotas_de_download_do_relatorio():
    app = create_app()
    search._SEARCHES["export00000001"] = _resultado_base()
    client = app.test_client()

    pdf = client.get("/relatorio/export00000001/pdf")
    assert pdf.status_code == 200
    assert pdf.headers["Content-Type"] == "application/pdf"
    assert "attachment" in pdf.headers["Content-Disposition"]
    assert pdf.data.startswith(b"%PDF")

    xlsx = client.get("/relatorio/export00000001/xlsx")
    assert xlsx.status_code == 200
    assert "spreadsheetml" in xlsx.headers["Content-Type"]
    assert xlsx.data.startswith(b"PK")


def test_rotas_de_download_do_comparativo(monkeypatch):
    cmp_id, _ = _comparativo(monkeypatch)
    app = create_app()
    client = app.test_client()

    pdf = client.get(f"/comparativo/{cmp_id}/pdf")
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF")

    xlsx = client.get(f"/comparativo/{cmp_id}/xlsx")
    assert xlsx.status_code == 200
    assert xlsx.data.startswith(b"PK")


def test_download_inexistente_retorna_404():
    app = create_app()
    client = app.test_client()
    assert client.get("/relatorio/naoexiste/pdf").status_code == 404
    assert client.get("/comparativo/naoexiste/xlsx").status_code == 404


def test_paginas_oferecem_botoes_de_download(monkeypatch):
    app = create_app()
    search._SEARCHES["export00000001"] = _resultado_base()
    client = app.test_client()

    html = client.get("/relatorio/export00000001").data.decode("utf-8")
    assert "/relatorio/export00000001/pdf" in html
    assert "/relatorio/export00000001/xlsx" in html

    cmp_id, _ = _comparativo(monkeypatch)
    html = client.get(f"/comparativo/{cmp_id}").data.decode("utf-8")
    assert f"/comparativo/{cmp_id}/pdf" in html
    assert f"/comparativo/{cmp_id}/xlsx" in html
