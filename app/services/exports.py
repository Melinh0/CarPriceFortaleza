from __future__ import annotations

import io
import time
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AZUL = colors.HexColor("#123a6b")
AZUL_ESCURO = colors.HexColor("#0d2c53")
CINZA = colors.HexColor("#66707d")
BORDA = colors.HexColor("#e2e7ee")
LINHA_ALTERNADA = colors.HexColor("#f7f9fc")
TEXTO = colors.HexColor("#1d2733")

CABECALHO_XLSX = "123A6B"
ROTULO_XLSX = "1F6FEB"
FORMATO_MOEDA = '"R$" #,##0.00'

SUBSTITUICOES = {
    "—": "-",
    "–": "-",
    "‘": "'",
    "’": "'",
    "“": '"',
    "”": '"',
    "·": "-",
    "→": "->",
    "≈": "~",
    "…": "...",
    "\u00a0": " ",
    " ": " ",
}


def _limpo(valor) -> str:
    if valor is None:
        return "-"
    texto = str(valor)
    for antigo, novo in SUBSTITUICOES.items():
        texto = texto.replace(antigo, novo)
    return texto.encode("cp1252", "replace").decode("cp1252")


def _texto(valor) -> str:
    return escape(_limpo(valor))


def _brl(valor) -> str:
    if valor is None:
        return "-"
    try:
        formato = f"{float(valor):,.2f}"
    except (TypeError, ValueError):
        return "-"
    return "R$ " + formato.replace(",", "X").replace(".", ",").replace("X", ".")


def _inteiro(valor) -> str:
    try:
        return f"{int(valor):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "-"


def _percent(valor, casas: int = 1) -> str:
    if valor is None:
        return "-"
    try:
        formato = f"{round(float(valor), casas):.{casas}f}"
    except (TypeError, ValueError):
        return "-"
    if casas:
        formato = formato.rstrip("0").rstrip(".")
    return formato.replace(".", ",")


def _moeda(valor):
    if valor is None:
        return None
    try:
        return round(float(valor), 2)
    except (TypeError, ValueError):
        return None


def _data(ts) -> str:
    try:
        return time.strftime("%d/%m/%Y %H:%M", time.localtime(float(ts)))
    except (TypeError, ValueError, OSError):
        return "-"


def _nome_relatorio(rel: dict) -> str:
    partes = f"{rel.get('marca') or ''} {rel.get('modelo') or ''}".strip()
    return partes or "Carros"


def _estilos_pdf() -> dict:
    return {
        "titulo": ParagraphStyle(
            "titulo",
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=21,
            textColor=AZUL,
        ),
        "sub": ParagraphStyle(
            "sub",
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=CINZA,
            spaceAfter=6,
        ),
        "secao": ParagraphStyle(
            "secao",
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=AZUL,
            spaceBefore=14,
            spaceAfter=6,
        ),
        "subsecao": ParagraphStyle(
            "subsecao",
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=12,
            textColor=AZUL,
            spaceBefore=6,
            spaceAfter=3,
        ),
        "corpo": ParagraphStyle(
            "corpo",
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=TEXTO,
            spaceAfter=5,
        ),
        "nota": ParagraphStyle(
            "nota",
            fontName="Helvetica-Oblique",
            fontSize=8.2,
            leading=11,
            textColor=CINZA,
            spaceBefore=10,
        ),
        "celula": ParagraphStyle(
            "celula",
            fontName="Helvetica",
            fontSize=7.6,
            leading=9.4,
            textColor=TEXTO,
        ),
        "celula_dir": ParagraphStyle(
            "celula_dir",
            fontName="Helvetica",
            fontSize=7.6,
            leading=9.4,
            textColor=TEXTO,
            alignment=2,
        ),
        "celula_rotulo": ParagraphStyle(
            "celula_rotulo",
            fontName="Helvetica-Bold",
            fontSize=7.8,
            leading=9.6,
            textColor=AZUL,
        ),
        "cabecalho": ParagraphStyle(
            "cabecalho",
            fontName="Helvetica-Bold",
            fontSize=7.6,
            leading=9.4,
            textColor=colors.white,
        ),
        "cabecalho_dir": ParagraphStyle(
            "cabecalho_dir",
            fontName="Helvetica-Bold",
            fontSize=7.6,
            leading=9.4,
            textColor=colors.white,
            alignment=2,
        ),
        "item": ParagraphStyle(
            "item",
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=TEXTO,
            leftIndent=12,
            bulletIndent=2,
            spaceAfter=3,
        ),
    }


def _tabela_pdf(
    linhas,
    larguras,
    estilos,
    colunas_num=(),
    colunas_rotulo=(),
    cabecalho=True,
):
    dados = []
    for i, linha in enumerate(linhas):
        celulas = []
        for j, celula in enumerate(linha):
            if isinstance(celula, Paragraph):
                celulas.append(celula)
                continue
            if i == 0 and cabecalho:
                estilo = estilos["cabecalho_dir" if j in colunas_num else "cabecalho"]
            elif j in colunas_rotulo:
                estilo = estilos["celula_rotulo"]
            else:
                estilo = estilos["celula_dir" if j in colunas_num else "celula"]
            celulas.append(Paragraph(_texto(celula), estilo))
        dados.append(celulas)

    tabela = Table(dados, colWidths=larguras, repeatRows=1 if cabecalho else 0)
    comandos = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDA),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if cabecalho:
        comandos += [
            ("BACKGROUND", (0, 0), (-1, 0), AZUL),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, AZUL_ESCURO),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LINHA_ALTERNADA]),
        ]
    else:
        comandos += [
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LINHA_ALTERNADA]),
        ]
        for coluna in colunas_rotulo:
            comandos.append(("BACKGROUND", (coluna, 0), (coluna, -1), colors.HexColor("#eef3fa")))
    tabela.setStyle(TableStyle(comandos))
    return tabela


def _add(elementos, estilos, titulo):
    elementos.append(Paragraph(_texto(titulo), estilos["secao"]))


def _add_tabela(elementos, linhas, larguras, estilos, **kwargs):
    elementos.append(_tabela_pdf(linhas, larguras, estilos, **kwargs))
    elementos.append(Spacer(1, 6))


def _add_itens(elementos, estilos, itens):
    for item in itens or []:
        elementos.append(Paragraph(_texto(item), estilos["item"], bulletText="•"))


def _duas_linhas(primeira, segunda) -> str:
    return f"{_texto(primeira)}<br/>{_texto(segunda)}"


def _rodape_pdf(titulo: str):
    def desenhar(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(CINZA)
        canvas.drawString(16 * mm, 10 * mm, _limpo(titulo))
        canvas.drawRightString(A4[0] - 16 * mm, 10 * mm, f"Página {doc.page}")
        canvas.restoreState()

    return desenhar


def _documento_pdf(titulo: str):
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=15 * mm,
        bottomMargin=17 * mm,
        title=_limpo(titulo),
        author="CarPrice Fortaleza",
    )
    return buffer, documento


def _avisos_relatorio(rel: dict) -> list[str]:
    avisos = []
    if not rel.get("fontes_ao_vivo"):
        avisos.append(
            "Origem dos dados: sem ofertas ao vivo nesta consulta — os preços abaixo são "
            "estimativas da base local de referência."
        )
    relaxados = rel.get("filtros_relaxados") or []
    if relaxados:
        avisos.append(
            f"Filtros relaxados: {', '.join(str(r) for r in relaxados)} — ampliados para o "
            "relatório não ficar incompleto."
        )
    total = (rel.get("stats") or {}).get("total", 0)
    if 0 < total < 3:
        avisos.append(
            f"Apenas {total} oferta(s) passaram nos filtros — amplie preço, ano ou km para "
            "enriquecer a comparação."
        )
    if total == 0:
        avisos.append(
            "Base local sem cobertura para esta busca — refaça escolhendo uma marca/modelo do "
            "catálogo."
        )
    return avisos


def relatorio_pdf(rel: dict) -> bytes:
    titulo_doc = f"Relatório: {_nome_relatorio(rel)}"
    buffer, documento = _documento_pdf(titulo_doc)
    estilos = _estilos_pdf()
    largura = A4[0] - 32 * mm
    elementos = []

    elementos.append(Paragraph(_texto(titulo_doc), estilos["titulo"]))
    elementos.append(
        Paragraph(
            _texto(
                "CarPrice Fortaleza · gerado em "
                f"{_data(rel.get('gerado_em'))} · relatório {rel.get('id') or '-'}"
            ),
            estilos["sub"],
        )
    )
    for aviso in _avisos_relatorio(rel):
        elementos.append(Paragraph(_texto(aviso), estilos["nota"]))

    stats = rel.get("stats") or {}
    _add(elementos, estilos, "Resumo executivo")
    pares = [
        ("Menor preço", _brl(stats.get("min"))),
        ("Mediana", _brl(stats.get("mediana"))),
        ("Média", _brl(stats.get("media"))),
        ("Maior preço", _brl(stats.get("max"))),
        ("Desvio padrão", _brl(stats.get("desvio"))),
        ("Ofertas analisadas", str(stats.get("total", 0))),
        ("Ofertas ao vivo", str(stats.get("ao_vivo", 0))),
    ]
    linhas_pares = []
    for i in range(0, len(pares), 2):
        esquerda = pares[i]
        direita = pares[i + 1] if i + 1 < len(pares) else ("", "")
        linhas_pares.append([esquerda[0], esquerda[1], direita[0], direita[1]])
    _add_tabela(
        elementos,
        linhas_pares,
        [largura * 0.24, largura * 0.26, largura * 0.24, largura * 0.26],
        estilos,
        colunas_num=(1, 3),
        colunas_rotulo=(0, 2),
        cabecalho=False,
    )

    melhor = rel.get("melhor_oferta")
    if melhor:
        elementos.append(
            Paragraph(
                _texto(
                    f"Melhor oferta: {melhor.get('titulo')} — {_brl(melhor.get('preco'))} "
                    f"({melhor.get('cidade') or 'Fortaleza - CE'})."
                ),
                estilos["corpo"],
            )
        )
    exemplo = rel.get("exemplo_financiamento")
    if exemplo:
        elementos.append(
            Paragraph(
                _texto(
                    f"Exemplo de financiamento da melhor oferta: entrada de "
                    f"{_brl(exemplo.get('entrada'))} ({_percent(exemplo.get('entrada_pct'))}), "
                    f"{exemplo.get('meses')} meses a {_percent(exemplo.get('taxa_aa'))}% a.a. — parcela de "
                    f"{_brl(exemplo.get('parcela'))}, total {_brl(exemplo.get('total_pago'))} "
                    f"(juros de {_brl(exemplo.get('juros'))})."
                ),
                estilos["corpo"],
            )
        )

    if rel.get("histograma"):
        _add(elementos, estilos, "Distribuição de preços")
        linhas = [["Faixa", "Ofertas", "%"]]
        for faixa in rel["histograma"]:
            linhas.append(
                [
                    faixa.get("faixa"),
                    str(faixa.get("total", 0)),
                    f"{faixa.get('percentual', 0)}%",
                ]
            )
        _add_tabela(
            elementos,
            linhas,
            [largura * 0.6, largura * 0.2, largura * 0.2],
            estilos,
            colunas_num=(1, 2),
        )

    _add(elementos, estilos, "Ofertas e preços")
    ofertas = rel.get("ofertas") or []
    if not ofertas:
        elementos.append(
            Paragraph(_texto("Nenhuma oferta nesta busca."), estilos["corpo"])
        )
    else:
        linhas = [["Veículo", "Ano", "Km", "Vendedor", "Preço", "Origem"]]
        for oferta in ofertas:
            linhas.append(
                [
                    oferta.get("titulo"),
                    str(oferta.get("ano") or "-"),
                    _inteiro(oferta.get("km")) if oferta.get("km") else "-",
                    oferta.get("vendedor") or "-",
                    _brl(oferta.get("preco")),
                    "ao vivo" if oferta.get("ao_vivo") else "estimativa",
                ]
            )
        _add_tabela(
            elementos,
            linhas,
            [
                largura * 0.34,
                largura * 0.09,
                largura * 0.12,
                largura * 0.19,
                largura * 0.15,
                largura * 0.11,
            ],
            estilos,
            colunas_num=(1, 2, 4),
        )

    if rel.get("comparativo"):
        _add(elementos, estilos, "Comparativo da marca")
        linhas = [["Modelo", "Combustível", "Carroceria", "Preço 0 km", "Usado ref."]]
        for item in rel["comparativo"]:
            linhas.append(
                [
                    item.get("modelo"),
                    item.get("combustivel"),
                    item.get("carroceria"),
                    _brl(item.get("preco_novo_ref")),
                    _brl(item.get("preco_usado_ref")),
                ]
            )
        _add_tabela(
            elementos,
            linhas,
            [
                largura * 0.26,
                largura * 0.18,
                largura * 0.18,
                largura * 0.19,
                largura * 0.19,
            ],
            estilos,
            colunas_num=(3, 4),
        )

    comentarios = rel.get("comentarios") or {}
    resumo = comentarios.get("resumo_base") or {}
    secoes_comentarios = [
        ("Pontos positivos relatados", resumo.get("pontos_positivos")),
        ("Pontos negativos relatados", resumo.get("pontos_negativos")),
        ("Sobre preço", resumo.get("sobre_preco")),
    ]
    web_comentarios = comentarios.get("comentarios_web") or []
    if any(itens for _, itens in secoes_comentarios) or web_comentarios:
        _add(elementos, estilos, "Comentários de usuários")
        for rotulo, itens in secoes_comentarios:
            if not itens:
                continue
            elementos.append(Paragraph(_texto(rotulo), estilos["subsecao"]))
            _add_itens(elementos, estilos, itens)
        if web_comentarios:
            elementos.append(Paragraph(_texto("O que a web diz"), estilos["subsecao"]))
            _add_itens(
                elementos,
                estilos,
                [
                    f"{c.get('titulo')} — {c.get('trecho')}" if c.get("trecho") else c.get("titulo")
                    for c in web_comentarios[:6]
                ],
            )

    if rel.get("mencoes_web"):
        _add(elementos, estilos, "Preços citados na web")
        _add_itens(
            elementos,
            estilos,
            [
                f"{m.get('titulo')} — {_brl(m.get('preco'))}"
                if m.get("preco")
                else str(m.get("titulo"))
                for m in rel["mencoes_web"][:8]
            ],
        )

    if rel.get("concessionarias"):
        _add(elementos, estilos, "Onde ir em Fortaleza — concessionárias")
        linhas = [["Concessionária", "Endereço", "Telefone", "Horário"]]
        for d in rel["concessionarias"]:
            linhas.append(
                [
                    d.get("nome"),
                    f"{d.get('endereco')} — {d.get('bairro')}",
                    d.get("telefone") or "-",
                    d.get("horario") or "-",
                ]
            )
        _add_tabela(
            elementos,
            linhas,
            [largura * 0.27, largura * 0.33, largura * 0.15, largura * 0.25],
            estilos,
        )

    if rel.get("dicas"):
        _add(elementos, estilos, "Dicas de negociação")
        _add_itens(elementos, estilos, rel["dicas"])

    elementos.append(
        Paragraph(
            _texto(
                "Relatório gerado pelo CarPrice Fortaleza. Preços estimados são referência de "
                "negociação e não substituem a proposta formal da concessionária."
            ),
            estilos["nota"],
        )
    )

    rodape = _rodape_pdf(f"CarPrice Fortaleza — relatório {rel.get('id') or ''}")
    documento.build(elementos, onFirstPage=rodape, onLaterPages=rodape)
    return buffer.getvalue()


def _linha_cabecalho(planilha, linha: int, textos) -> int:
    preenchimento = PatternFill("solid", fgColor=CABECALHO_XLSX)
    for coluna, texto in enumerate(textos, start=1):
        celula = planilha.cell(row=linha, column=coluna, value=texto)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = preenchimento
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    planilha.row_dimensions[linha].height = 24
    return linha + 1


def _larguras(planilha, larguras) -> None:
    for coluna, largura in enumerate(larguras, start=1):
        planilha.column_dimensions[get_column_letter(coluna)].width = largura


def _celula_moeda(planilha, linha: int, coluna: int, valor) -> None:
    celula = planilha.cell(row=linha, column=coluna, value=_moeda(valor))
    celula.number_format = FORMATO_MOEDA


def _celula_link(planilha, linha: int, coluna: int, url) -> None:
    if not url:
        return
    celula = planilha.cell(row=linha, column=coluna, value=url)
    celula.hyperlink = url
    celula.font = Font(color=ROTULO_XLSX, underline="single")


def _rotulo(planilha, linha: int, texto: str) -> None:
    celula = planilha.cell(row=linha, column=1, value=texto)
    celula.font = Font(bold=True, color=ROTULO_XLSX)


def _resumo_filtros(filtros: dict) -> str:
    partes = []
    if filtros.get("marca"):
        partes.append(f"marca {filtros['marca']}")
    if filtros.get("modelo"):
        partes.append(f"modelo {filtros['modelo']}")
    if filtros.get("preco_min"):
        partes.append(f"preço mín. {_brl(filtros['preco_min'])}")
    if filtros.get("preco_max"):
        partes.append(f"preço máx. {_brl(filtros['preco_max'])}")
    if filtros.get("ano_min"):
        partes.append(f"ano mín. {filtros['ano_min']}")
    if filtros.get("ano_max"):
        partes.append(f"ano máx. {filtros['ano_max']}")
    if filtros.get("km_max"):
        partes.append(f"km até {_inteiro(filtros['km_max'])}")
    if filtros.get("combustivel"):
        partes.append(f"combustível {filtros['combustivel']}")
    if filtros.get("cambio"):
        partes.append(f"câmbio {filtros['cambio']}")
    if filtros.get("carroceria"):
        partes.append(f"carroceria {filtros['carroceria']}")
    if filtros.get("caracteristicas"):
        partes.append(", ".join(filtros["caracteristicas"]))
    return " · ".join(partes) or "sem filtros"


def _aba_resumo(planilha, rel: dict) -> None:
    _larguras(planilha, [26, 62, 16, 14])
    linha = 1
    identificacao = [
        ("Relatório", _nome_relatorio(rel)),
        ("Gerado em", _data(rel.get("gerado_em"))),
        ("ID", str(rel.get("id") or "-")),
        ("Filtros", _resumo_filtros(rel.get("filtros") or {})),
    ]
    for rotulo, valor in identificacao:
        _rotulo(planilha, linha, rotulo)
        planilha.cell(row=linha, column=2, value=valor)
        linha += 1

    linha += 1
    _rotulo(planilha, linha, "Estatísticas")
    linha += 1
    linha = _linha_cabecalho(planilha, linha, ["Indicador", "Valor"])
    stats = rel.get("stats") or {}
    indicadores = [
        ("Menor preço", stats.get("min"), True),
        ("Mediana", stats.get("mediana"), True),
        ("Média", stats.get("media"), True),
        ("Maior preço", stats.get("max"), True),
        ("Desvio padrão", stats.get("desvio"), True),
        ("Ofertas analisadas", stats.get("total", 0), False),
        ("Ofertas ao vivo", stats.get("ao_vivo", 0), False),
    ]
    for rotulo, valor, moeda in indicadores:
        planilha.cell(row=linha, column=1, value=rotulo)
        if moeda:
            _celula_moeda(planilha, linha, 2, valor)
        else:
            planilha.cell(row=linha, column=2, value=valor)
        linha += 1

    if rel.get("histograma"):
        linha += 1
        _rotulo(planilha, linha, "Distribuição de preços")
        linha += 1
        linha = _linha_cabecalho(planilha, linha, ["Faixa", "Ofertas", "%"])
        for faixa in rel["histograma"]:
            planilha.cell(row=linha, column=1, value=faixa.get("faixa"))
            planilha.cell(row=linha, column=2, value=faixa.get("total", 0))
            planilha.cell(row=linha, column=3, value=faixa.get("percentual", 0))
            linha += 1

    melhor = rel.get("melhor_oferta")
    if melhor:
        linha += 1
        _rotulo(planilha, linha, "Melhor oferta")
        planilha.cell(
            row=linha,
            column=2,
            value=f"{melhor.get('titulo')} — {_brl(melhor.get('preco'))}",
        )
        linha += 1

    linha += 1
    nota = planilha.cell(
        row=linha,
        column=1,
        value=(
            "Gerado pelo CarPrice Fortaleza. Preços estimados são referência de negociação e "
            "não substituem a proposta formal da concessionária."
        ),
    )
    nota.font = Font(italic=True, color="66707D")
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    planilha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=4)
    planilha.row_dimensions[linha].height = 30


def _aba_ofertas(planilha, rel: dict) -> None:
    _larguras(planilha, [46, 8, 12, 40, 18, 16, 15, 12, 60])
    _linha_cabecalho(
        planilha,
        1,
        [
            "Veículo",
            "Ano",
            "Km",
            "Características",
            "Vendedor",
            "Cidade",
            "Preço",
            "Origem",
            "Link",
        ],
    )
    for linha, oferta in enumerate(rel.get("ofertas") or [], start=2):
        planilha.cell(row=linha, column=1, value=oferta.get("titulo"))
        planilha.cell(row=linha, column=2, value=oferta.get("ano"))
        planilha.cell(row=linha, column=3, value=oferta.get("km"))
        planilha.cell(
            row=linha,
            column=4,
            value=", ".join(oferta.get("caracteristicas") or []),
        )
        planilha.cell(row=linha, column=5, value=oferta.get("vendedor"))
        planilha.cell(row=linha, column=6, value=oferta.get("cidade"))
        _celula_moeda(planilha, linha, 7, oferta.get("preco"))
        planilha.cell(
            row=linha,
            column=8,
            value="ao vivo" if oferta.get("ao_vivo") else "estimativa",
        )
        _celula_link(planilha, linha, 9, oferta.get("url"))
    planilha.auto_filter.ref = planilha.dimensions
    planilha.freeze_panes = "A2"


def _aba_concessionarias(planilha, rel: dict) -> None:
    _larguras(planilha, [34, 46, 20, 16, 20, 34, 56])
    _linha_cabecalho(
        planilha,
        1,
        ["Concessionária", "Endereço", "Bairro", "Cidade/UF", "Telefone", "Horário", "Rota"],
    )
    for linha, d in enumerate(rel.get("concessionarias") or [], start=2):
        planilha.cell(row=linha, column=1, value=d.get("nome"))
        planilha.cell(row=linha, column=2, value=d.get("endereco"))
        planilha.cell(row=linha, column=3, value=d.get("bairro"))
        planilha.cell(
            row=linha,
            column=4,
            value=f"{d.get('cidade')}/{d.get('uf')}",
        )
        planilha.cell(row=linha, column=5, value=d.get("telefone"))
        planilha.cell(row=linha, column=6, value=d.get("horario"))
        _celula_link(planilha, linha, 7, d.get("rota_url"))


def _aba_comparativo_marca(planilha, rel: dict) -> None:
    _larguras(planilha, [30, 18, 18, 18, 18])
    _linha_cabecalho(
        planilha, 1, ["Modelo", "Combustível", "Carroceria", "Preço 0 km", "Usado ref."]
    )
    for linha, item in enumerate(rel.get("comparativo") or [], start=2):
        planilha.cell(row=linha, column=1, value=item.get("modelo"))
        planilha.cell(row=linha, column=2, value=item.get("combustivel"))
        planilha.cell(row=linha, column=3, value=item.get("carroceria"))
        _celula_moeda(planilha, linha, 4, item.get("preco_novo_ref"))
        _celula_moeda(planilha, linha, 5, item.get("preco_usado_ref"))


def _aba_comentarios(planilha, rel: dict) -> None:
    _larguras(planilha, [34, 74, 52])
    comentarios = rel.get("comentarios") or {}
    resumo = comentarios.get("resumo_base") or {}
    linha = 1
    for rotulo, chave in (
        ("Pontos positivos", "pontos_positivos"),
        ("Pontos negativos", "pontos_negativos"),
        ("Sobre preço", "sobre_preco"),
    ):
        itens = resumo.get(chave) or []
        if not itens:
            continue
        _rotulo(planilha, linha, rotulo)
        celula = planilha.cell(row=linha, column=2, value="; ".join(str(i) for i in itens))
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        planilha.row_dimensions[linha].height = 30
        linha += 1

    web = comentarios.get("comentarios_web") or []
    if web:
        linha += 1
        linha = _linha_cabecalho(planilha, linha, ["Comentários da web", "Trecho", "Link"])
        for item in web:
            planilha.cell(row=linha, column=1, value=item.get("titulo"))
            celula = planilha.cell(row=linha, column=2, value=item.get("trecho"))
            celula.alignment = Alignment(wrap_text=True, vertical="top")
            _celula_link(planilha, linha, 3, item.get("url"))
            linha += 1


def _aba_mencoes(planilha, rel: dict) -> None:
    _larguras(planilha, [56, 16, 74, 60])
    _linha_cabecalho(planilha, 1, ["Título", "Preço", "Trecho", "Link"])
    for linha, m in enumerate(rel.get("mencoes_web") or [], start=2):
        planilha.cell(row=linha, column=1, value=m.get("titulo"))
        _celula_moeda(planilha, linha, 2, m.get("preco"))
        celula = planilha.cell(row=linha, column=3, value=m.get("trecho"))
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        _celula_link(planilha, linha, 4, m.get("url"))


def relatorio_xlsx(rel: dict) -> bytes:
    pasta = Workbook()
    _aba_resumo(pasta.active, rel)
    pasta.active.title = "Resumo"
    _aba_ofertas(pasta.create_sheet("Ofertas"), rel)
    _aba_concessionarias(pasta.create_sheet("Concessionárias"), rel)
    _aba_comparativo_marca(pasta.create_sheet("Comparativo da marca"), rel)
    _aba_comentarios(pasta.create_sheet("Comentários"), rel)
    _aba_mencoes(pasta.create_sheet("Menções web"), rel)
    buffer = io.BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()


def _entrada_txt(condicoes: dict) -> str:
    if condicoes.get("entrada") is not None:
        return _brl(condicoes.get("entrada"))
    return f"{_percent(condicoes.get('entrada_pct'))}% do valor de cada carro"


def _garantia_txt(garantia: dict, chave_anos: str, chave_km: str) -> str:
    anos = garantia.get(chave_anos)
    if not anos:
        return "-"
    km = garantia.get(chave_km)
    if km:
        return f"{anos} ano(s) ou {_inteiro(km)} km"
    return f"{anos} ano(s) sem limite de km"


def comparativo_pdf(rel: dict) -> bytes:
    buffer, documento = _documento_pdf("Comparativo de carros novos")
    estilos = _estilos_pdf()
    largura = A4[0] - 32 * mm
    elementos = []
    cond = rel.get("condicoes") or {}

    elementos.append(
        Paragraph(_texto("Comparativo de carros novos"), estilos["titulo"])
    )
    elementos.append(
        Paragraph(
            _texto(
                "CarPrice Fortaleza · gerado em "
                f"{rel.get('gerado_em') or _data(rel.get('ts'))} · relatório "
                f"{rel.get('id') or '-'}"
            ),
            estilos["sub"],
        )
    )
    elementos.append(
        Paragraph(
            _texto(
                f"Condições: entrada {_entrada_txt(cond)} · juros "
                f"{_percent(cond.get('taxa_aa'))}% a.a. em {cond.get('meses')}x · sem juros "
                f"{cond.get('meses_sem_juros')}x · consórcio {cond.get('meses_consorcio')}x "
                f"(adm. {_percent(cond.get('taxa_adm'))}%)"
            ),
            estilos["nota"],
        )
    )

    _add(elementos, estilos, "Preços e garantia — lado a lado")
    linhas = [
        [
            "Veículo",
            "Preço tabela",
            "Preço à vista",
            "Economia",
            "Garantia do veículo",
            "Garantia da bateria",
        ]
    ]
    for c in rel.get("carros") or []:
        garantia = c.get("garantia") or {}
        linhas.append(
            [
                f"{c.get('marca')} {c.get('modelo')}",
                _brl(c.get("preco_loja")),
                _brl(c.get("preco_a_vista")),
                f"{_brl(c.get('economia'))} ({_percent(c.get('economia_pct'))}%)",
                _garantia_txt(garantia, "anos", "km"),
                _garantia_txt(garantia, "bateria_anos", "bateria_km"),
            ]
        )
    _add_tabela(
        elementos,
        linhas,
        [
            largura * 0.22,
            largura * 0.15,
            largura * 0.15,
            largura * 0.16,
            largura * 0.16,
            largura * 0.16,
        ],
        estilos,
        colunas_num=(1, 2, 3),
    )

    _add(
        elementos,
        estilos,
        f"Parcelas — entrada de {_entrada_txt(cond)}",
    )
    linhas = [["Veículo", "À vista", "Com juros", "Sem juros", "Consórcio"]]
    for c in rel.get("carros") or []:
        cenarios = c.get("cenarios") or {}
        cj = cenarios.get("com_juros") or {}
        sj = cenarios.get("sem_juros") or {}
        co = cenarios.get("consorcio") or {}
        av = cenarios.get("a_vista") or {}
        linhas.append(
            [
                f"{c.get('marca')} {c.get('modelo')}",
                Paragraph(_texto(_brl(av.get("total_pago"))), estilos["celula"]),
                Paragraph(
                    _duas_linhas(
                        f"{_brl(cj.get('parcela'))}/mês",
                        f"total {_brl(cj.get('total_pago'))} · juros {_brl(cj.get('juros'))}",
                    ),
                    estilos["celula"],
                ),
                Paragraph(
                    _duas_linhas(
                        f"{_brl(sj.get('parcela'))}/mês",
                        f"total {_brl(sj.get('total_pago'))} · juros {_brl(sj.get('juros'))}",
                    ),
                    estilos["celula"],
                ),
                Paragraph(
                    _duas_linhas(
                        f"{_brl(co.get('parcela'))}/mês",
                        f"total {_brl(co.get('total_pago'))} · adm. "
                        f"{_percent(co.get('taxa_adm_pct'))}%",
                    ),
                    estilos["celula"],
                ),
            ]
        )
    _add_tabela(
        elementos,
        linhas,
        [
            largura * 0.24,
            largura * 0.16,
            largura * 0.2,
            largura * 0.2,
            largura * 0.2,
        ],
        estilos,
    )

    for c in rel.get("carros") or []:
        cenarios = c.get("cenarios") or {}
        cj = cenarios.get("com_juros") or {}
        sj = cenarios.get("sem_juros") or {}
        co = cenarios.get("consorcio") or {}
        garantia = c.get("garantia") or {}
        _add(elementos, estilos, f"{c.get('marca')} {c.get('modelo')}")
        linhas = [["Campo", "Valor"]]
        linhas += [
            ["Preço de tabela / loja virtual", _brl(c.get("preco_loja"))],
            ["Preço à vista de negociação", _brl(c.get("preco_a_vista"))],
            [
                "Economia frente ao anunciado",
                f"{_brl(c.get('economia'))} ({_percent(c.get('economia_pct'))}%)",
            ],
            [
                "Origem do preço à vista",
                "web ao vivo" if c.get("origem_avista") == "web" else "estimativa local",
            ],
            ["Ficha", f"{c.get('motor')} · {c.get('cambio')} · {c.get('carroceria')} · {c.get('combustivel')}"],
            ["Garantia do veículo", _garantia_txt(garantia, "anos", "km")],
            ["Garantia da bateria", _garantia_txt(garantia, "bateria_anos", "bateria_km")],
            ["Entrada", f"{_brl(cenarios.get('entrada'))} ({_percent(cenarios.get('entrada_pct'))})"],
            ["Valor financiado", _brl(cj.get("financiado"))],
            ["Taxa com juros", f"{_percent(cj.get('taxa_aa'))}% a.a. em {cj.get('meses')}x"],
            ["Parcela com juros", _brl(cj.get("parcela"))],
            ["Total com juros", _brl(cj.get("total_pago"))],
            ["Juros do crédito", _brl(cj.get("juros"))],
            ["Parcela sem juros", f"{_brl(sj.get('parcela'))} em {sj.get('meses')}x"],
            ["Total sem juros", _brl(sj.get("total_pago"))],
            [
                "Consórcio",
                f"{_brl(co.get('parcela'))} em {co.get('meses')}x · total {_brl(co.get('total_pago'))}",
            ],
        ]
        _add_tabela(
            elementos,
            linhas,
            [largura * 0.42, largura * 0.58],
            estilos,
            colunas_rotulo=(0,),
        )

    lojas = {}
    for c in rel.get("carros") or []:
        for d in c.get("concessionarias") or []:
            lojas.setdefault(d.get("nome"), d)
    if lojas:
        _add(elementos, estilos, "Onde negociar em Fortaleza")
        linhas = [["Concessionária", "Marcas", "Endereço", "Telefone"]]
        for d in lojas.values():
            linhas.append(
                [
                    d.get("nome"),
                    ", ".join(d.get("marcas") or []),
                    f"{d.get('endereco')} — {d.get('bairro')}",
                    d.get("telefone") or "-",
                ]
            )
        _add_tabela(
            elementos,
            linhas,
            [largura * 0.26, largura * 0.22, largura * 0.34, largura * 0.18],
            estilos,
        )

    if rel.get("dicas"):
        _add(elementos, estilos, "Dicas de negociação")
        _add_itens(elementos, estilos, rel["dicas"])

    elementos.append(
        Paragraph(
            _texto(
                "Comparativo gerado pelo CarPrice Fortaleza. Preços de tabela e garantias são "
                "dados de referência — o valor à vista definitivo é o da proposta escrita da "
                "concessionária. Simulações sem IOF e tarifas: confirme o CET no contrato."
            ),
            estilos["nota"],
        )
    )

    rodape = _rodape_pdf(f"CarPrice Fortaleza — comparativo {rel.get('id') or ''}")
    documento.build(elementos, onFirstPage=rodape, onLaterPages=rodape)
    return buffer.getvalue()


def _aba_condicoes(planilha, rel: dict) -> None:
    _larguras(planilha, [30, 62])
    cond = rel.get("condicoes") or {}
    linha = 1
    identificacao = [
        ("Relatório", "Comparativo de carros novos"),
        ("Gerado em", rel.get("gerado_em") or _data(rel.get("ts"))),
        ("ID", str(rel.get("id") or "-")),
        ("Entrada", _entrada_txt(cond)),
        ("Juros", f"{_percent(cond.get('taxa_aa'))}% a.a. em {cond.get('meses')}x"),
        ("Sem juros", f"{cond.get('meses_sem_juros')}x"),
        (
            "Consórcio",
            f"{cond.get('meses_consorcio')}x · adm. {_percent(cond.get('taxa_adm'))}%",
        ),
        ("Carros", str(len(rel.get("carros") or []))),
    ]
    for rotulo, valor in identificacao:
        _rotulo(planilha, linha, rotulo)
        planilha.cell(row=linha, column=2, value=valor)
        linha += 1

    linha += 1
    _rotulo(planilha, linha, "Fontes")
    linha += 1
    linha = _linha_cabecalho(planilha, linha, ["Fonte", "Detalhe"])
    for fonte in rel.get("fontes") or []:
        planilha.cell(row=linha, column=1, value=fonte.get("nome"))
        planilha.cell(
            row=linha,
            column=2,
            value=f"{'ok' if fonte.get('ok') else 'parcial'} — {fonte.get('detalhe')}",
        )
        linha += 1


def _aba_precos(planilha, rel: dict) -> None:
    _larguras(planilha, [32, 17, 17, 17, 14, 26, 26, 20])
    _linha_cabecalho(
        planilha,
        1,
        [
            "Veículo",
            "Preço tabela",
            "Preço à vista",
            "Economia",
            "Economia %",
            "Garantia do veículo",
            "Garantia da bateria",
            "Origem do à vista",
        ],
    )
    for linha, c in enumerate(rel.get("carros") or [], start=2):
        garantia = c.get("garantia") or {}
        planilha.cell(row=linha, column=1, value=f"{c.get('marca')} {c.get('modelo')}")
        _celula_moeda(planilha, linha, 2, c.get("preco_loja"))
        _celula_moeda(planilha, linha, 3, c.get("preco_a_vista"))
        _celula_moeda(planilha, linha, 4, c.get("economia"))
        planilha.cell(
            row=linha, column=5, value=round(c.get("economia_pct") or 0, 1)
        )
        planilha.cell(
            row=linha, column=6, value=_garantia_txt(garantia, "anos", "km")
        )
        planilha.cell(
            row=linha,
            column=7,
            value=_garantia_txt(garantia, "bateria_anos", "bateria_km"),
        )
        planilha.cell(
            row=linha,
            column=8,
            value="web ao vivo" if c.get("origem_avista") == "web" else "estimativa local",
        )


def _aba_parcelas(planilha, rel: dict) -> None:
    _larguras(planilha, [32, 17, 17, 17, 17, 17, 17, 17, 17])
    _linha_cabecalho(
        planilha,
        1,
        [
            "Veículo",
            "À vista",
            "Com juros parcela",
            "Com juros total",
            "Com juros juros",
            "Sem juros parcela",
            "Sem juros total",
            "Consórcio parcela",
            "Consórcio total",
        ],
    )
    for linha, c in enumerate(rel.get("carros") or [], start=2):
        cenarios = c.get("cenarios") or {}
        cj = cenarios.get("com_juros") or {}
        sj = cenarios.get("sem_juros") or {}
        co = cenarios.get("consorcio") or {}
        av = cenarios.get("a_vista") or {}
        planilha.cell(row=linha, column=1, value=f"{c.get('marca')} {c.get('modelo')}")
        _celula_moeda(planilha, linha, 2, av.get("total_pago"))
        _celula_moeda(planilha, linha, 3, cj.get("parcela"))
        _celula_moeda(planilha, linha, 4, cj.get("total_pago"))
        _celula_moeda(planilha, linha, 5, cj.get("juros"))
        _celula_moeda(planilha, linha, 6, sj.get("parcela"))
        _celula_moeda(planilha, linha, 7, sj.get("total_pago"))
        _celula_moeda(planilha, linha, 8, co.get("parcela"))
        _celula_moeda(planilha, linha, 9, co.get("total_pago"))


def _aba_detalhes(planilha, rel: dict) -> None:
    _larguras(
        planilha,
        [32, 16, 16, 16, 14, 15, 15, 14, 12, 15, 15, 14, 16, 14],
    )
    _linha_cabecalho(
        planilha,
        1,
        [
            "Veículo",
            "Motor",
            "Câmbio",
            "Carroceria",
            "Combustível",
            "Entrada",
            "Financiado",
            "Taxa a.a.",
            "Prazo (meses)",
            "Parcela",
            "Total pago",
            "Juros",
            "Garantia (anos)",
            "Bateria (anos)",
        ],
    )
    for linha, c in enumerate(rel.get("carros") or [], start=2):
        cenarios = c.get("cenarios") or {}
        cj = cenarios.get("com_juros") or {}
        garantia = c.get("garantia") or {}
        planilha.cell(row=linha, column=1, value=f"{c.get('marca')} {c.get('modelo')}")
        planilha.cell(row=linha, column=2, value=c.get("motor"))
        planilha.cell(row=linha, column=3, value=c.get("cambio"))
        planilha.cell(row=linha, column=4, value=c.get("carroceria"))
        planilha.cell(row=linha, column=5, value=c.get("combustivel"))
        _celula_moeda(planilha, linha, 6, cenarios.get("entrada"))
        _celula_moeda(planilha, linha, 7, cj.get("financiado"))
        planilha.cell(row=linha, column=8, value=cj.get("taxa_aa"))
        planilha.cell(row=linha, column=9, value=cj.get("meses"))
        _celula_moeda(planilha, linha, 10, cj.get("parcela"))
        _celula_moeda(planilha, linha, 11, cj.get("total_pago"))
        _celula_moeda(planilha, linha, 12, cj.get("juros"))
        planilha.cell(row=linha, column=13, value=garantia.get("anos"))
        planilha.cell(row=linha, column=14, value=garantia.get("bateria_anos"))


def _aba_lojas(planilha, rel: dict) -> None:
    _larguras(planilha, [34, 26, 46, 20, 20, 34, 56])
    _linha_cabecalho(
        planilha,
        1,
        ["Concessionária", "Marcas", "Endereço", "Bairro", "Telefone", "Horário", "Rota"],
    )
    linhas = {}
    for c in rel.get("carros") or []:
        for d in c.get("concessionarias") or []:
            linhas.setdefault(d.get("nome"), d)
    for linha, d in enumerate(linhas.values(), start=2):
        planilha.cell(row=linha, column=1, value=d.get("nome"))
        planilha.cell(
            row=linha, column=2, value=", ".join(d.get("marcas") or [])
        )
        planilha.cell(row=linha, column=3, value=d.get("endereco"))
        planilha.cell(row=linha, column=4, value=d.get("bairro"))
        planilha.cell(row=linha, column=5, value=d.get("telefone"))
        planilha.cell(row=linha, column=6, value=d.get("horario"))
        _celula_link(planilha, linha, 7, d.get("rota_url"))


def comparativo_xlsx(rel: dict) -> bytes:
    pasta = Workbook()
    _aba_condicoes(pasta.active, rel)
    pasta.active.title = "Resumo"
    _aba_precos(pasta.create_sheet("Preços"), rel)
    _aba_parcelas(pasta.create_sheet("Parcelas"), rel)
    _aba_detalhes(pasta.create_sheet("Detalhes"), rel)
    _aba_lojas(pasta.create_sheet("Concessionárias"), rel)
    buffer = io.BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()
